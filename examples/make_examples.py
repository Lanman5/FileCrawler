"""Regenerate the example matrix and sample submission.

    python examples/make_examples.py                 # writes the .xlsx, .csv and .zip next to this file
    python examples/make_examples.py --folder DIR    # also unpacks the sample submission into DIR

The sample submission is designed to be checked against "Phase 2 - Design" and
hits every status: present, missing, empty folder, zero-byte file, wrong type,
misfiled, nested sub-folders, leading zeros ("02.09"), junk files and an
unrecognised folder.
"""
from __future__ import annotations

import argparse
import csv
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
PHASES = ["Phase 1 - Tender", "Phase 2 - Design", "Phase 3 - Construction", "Phase 4 - Handover"]

# (code, title, allowed types, required per phase as Y/N for the four phases)
MATRIX = [
    ("1", "General & Administration", "", ""),
    ("1.1", "Contractor Information Form", ".pdf; .docx", "YYYY"),
    ("1.2", "Insurance Certificates", ".pdf", "YYYY"),
    ("1.3", "Health & Safety Policy", ".pdf; .docx", "YYYN"),
    ("1.4", "Quality Management Plan", ".pdf; .docx", "YYYN"),
    ("2", "Design", "", ""),
    ("2.1", "Design Basis Report", ".pdf; .docx", "YYNN"),
    ("2.2", "Architectural Drawings", ".pdf; .dwg", "NYYN"),
    ("2.3", "Structural Calculations", ".pdf; .xlsx", "NYYN"),
    ("2.4", "Structural Drawings", ".pdf; .dwg", "NYYN"),
    ("2.5", "Mechanical Design Report", "", "NYNN"),
    ("2.6", "Electrical Design Report", "", "NYNN"),
    ("2.7", "Plumbing & Drainage Design", "", "NYNN"),
    ("2.8", "Fire Strategy Report", ".pdf", "NYYN"),
    ("2.9", "Geotechnical Report", ".pdf", "NYNN"),
    ("2.10", "Design Risk Register", ".xlsx; .pdf", "NYYN"),
    ("2.11", "Design Change Log", ".xlsx", "NYYY"),
    ("3", "Planning & Permits", "", ""),
    ("3.1", "Planning Application", ".pdf", "YYNN"),
    ("3.2", "Building Regulations Approval", ".pdf", "NYYN"),
    ("3.3", "Environmental Impact Assessment", ".pdf", "NYNN"),
    ("3.4", "Utility Connection Agreements", ".pdf", "NNYN"),
    ("3.5", "Road Closure Permits", ".pdf", "NNYN"),
    ("4", "Procurement", "", ""),
    ("4.1", "Procurement Strategy", "", "YYNN"),
    ("4.2", "Subcontractor List", ".xlsx; .pdf", "NYYN"),
    ("4.3", "Long-lead Items Schedule", ".xlsx", "NYYN"),
    ("4.4", "Material Submittals", "", "NNYN"),
    ("4.5", "Supplier Warranties", ".pdf", "NNYY"),
    ("4.6", "Purchase Order Log", ".xlsx", "NNYN"),
    ("5", "Construction", "", ""),
    ("5.1", "Construction Programme", ".pdf; .mpp; .xlsx", "YNYN"),
    ("5.2", "Method Statements", ".pdf; .docx", "NNYN"),
    ("5.3", "Site Logistics Plan", ".pdf", "NNYN"),
    ("5.4", "Daily Site Diaries", "", "NNYN"),
    ("5.5", "Monthly Progress Reports", ".pdf", "NNYN"),
    ("5.6", "Temporary Works Register", ".xlsx", "NNYN"),
    ("5.7", "Site Photographs", ".jpg; .png; .zip", "NNYY"),
    ("5.8", "Variations Register", ".xlsx", "NNYY"),
    ("6", "Health, Safety & Environment", "", ""),
    ("6.1", "Health & Safety Plan", ".pdf; .docx", "YYYN"),
    ("6.2", "Risk Assessments", "", "NYYN"),
    ("6.3", "COSHH Register", ".xlsx; .pdf", "NNYN"),
    ("6.4", "Environmental Management Plan", ".pdf", "NYYN"),
    ("6.5", "Site Waste Management Plan", ".pdf", "NNYN"),
    ("6.6", "Incident Log", ".xlsx", "NNYY"),
    ("6.7", "Site Induction & Training Records", "", "NNYN"),
    ("7", "Quality Assurance", "", ""),
    ("7.1", "Inspection & Test Plan", ".pdf; .xlsx", "YYYN"),
    ("7.2", "Material Test Certificates", ".pdf", "NNYY"),
    ("7.3", "Non-conformance Register", ".xlsx", "NNYY"),
    ("7.4", "Quality Audit Reports", ".pdf", "NNYN"),
    ("7.5", "Snagging List", ".xlsx; .pdf", "NNNY"),
    ("7.6", "Quality Records Index", ".xlsx", "NNNY"),
    ("8", "Testing & Commissioning", "", ""),
    ("8.1", "Commissioning Plan", ".pdf", "NNYN"),
    ("8.2", "Electrical Test Certificates", ".pdf", "NNNY"),
    ("8.3", "Mechanical Commissioning Records", ".pdf", "NNNY"),
    ("8.4", "Fire Alarm Certificate", ".pdf", "NNNY"),
    ("8.5", "Witness Test Reports", ".pdf", "NNNY"),
    ("9", "Handover & Close-out", "", ""),
    ("9.1", "As-built Drawings", ".pdf; .dwg", "NNNY"),
    ("9.2", "O&M Manuals", ".pdf", "NNNY"),
    ("9.3", "Health & Safety File", ".pdf", "NNNY"),
    ("9.4", "Building Log Book", ".pdf", "NNNY"),
    ("9.5", "Asset Register", ".xlsx", "NNNY"),
    ("9.6", "Warranties & Guarantees", ".pdf", "NNNY"),
    ("9.7", "End-user Training Records", ".pdf", "NNNY"),
    ("9.8", "Practical Completion Certificate", ".pdf", "NNNY"),
    ("9.9", "Final Account", ".pdf; .xlsx", "NNNY"),
]

ROOT = "ABC_Construction_Phase2_Submission"
# path relative to ROOT -> file content (None = empty folder, b"" = zero-byte file)
SUBMISSION = {
    "1 General & Administration/1.1 Contractor Information/Contractor_Information_Form.pdf": b"%PDF demo",
    "1 General & Administration/1.2 Insurance Certificates/Insurance_Certificate_2026.pdf": b"%PDF demo",
    "1 General & Administration/1.2 Insurance Certificates/Thumbs.db": b"junk",
    "1 General & Administration/1.3 Health and Safety Policy/": None,
    "1 General & Administration/1.4 Quality Management Plan/QMP_Rev2.docx": b"docx demo",
    "1 General & Administration/1.4 Quality Management Plan/~$QMP_Rev2.docx": b"lock",
    "2 Design/2.1 Design Basis Report/Design_Basis_Report_Rev1.pdf": b"%PDF demo",
    "2 Design/2.2 Architectural Drawings/Rev C/A-100_General_Arrangement.pdf": b"%PDF demo",
    "2 Design/2.2 Architectural Drawings/Rev C/A-100_General_Arrangement.dwg": b"dwg demo",
    "2 Design/2.3 Structural Calculations/Structural_Calcs_Rev0.xlsx": b"xlsx demo",
    "2 Design/2.4 Structural Drawings/S-100_markup_photo.jpg": b"jpg demo",
    "2 Design/2.6 Electrical Design Report/Electrical_Design_Report.docx": b"",
    "2 Design/2.7 Plumbing and Drainage/Drainage_Design.pdf": b"%PDF demo",
    "2 Design/2.8 Fire Strategy Report.pdf": b"%PDF demo",
    "2 Design/02.09 Geotechnical Report/Site_Investigation_Report.pdf": b"%PDF demo",
    "2 Design/2.10 Design Risk Register/Design_Risk_Register.xlsx": b"xlsx demo",
    "2 Design/2.11 Design Change Log/Design_Change_Log.xlsx": b"xlsx demo",
    "3 Planning & Permits/3.1 Planning Application/Planning_Application.pdf": b"%PDF demo",
    "3 Planning & Permits/3.3 Environmental Impact Assessment/EIA_Report.pdf": b"%PDF demo",
    "4 Procurement/4.1 Procurement Strategy/Procurement_Strategy.pdf": b"%PDF demo",
    "4 Procurement/4.2 Subcontractor List/Subcontractors.csv": b"name,trade",
    "4 Procurement/4.3 Long-lead Items/Long_Lead_Items.xlsx": b"xlsx demo",
    "5 Construction/5.1 Construction Programme/Outline_Programme.pdf": b"%PDF demo",
    "6 Health Safety & Environment/6.1 Health and Safety Plan/HS_Plan_Rev1.pdf": b"%PDF demo",
    "6 Health Safety & Environment/6.2 Risk Assessments/RA_Site_Survey.pdf": b"%PDF demo",
    "7 Quality/7.1 Inspection and Test Plan/ITP_Rev0.xlsx": b"xlsx demo",
    "7 Quality/7.9 Misc Correspondence/Letter_to_client.pdf": b"%PDF demo",
    "__MACOSX/2 Design/._Design_Basis_Report_Rev1.pdf": b"junk",
}


def write_xlsx(path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Matrix"
    sheet["A1"] = "Document Requirements Matrix - Example Project"
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A2"] = "Y = required in that phase. Keep the Code column formatted as Text (2.10 is not 2.1)."
    sheet["A2"].font = Font(italic=True, color="6C757D")
    header_row = 4
    header = ["Code", "Title", "Allowed Types"] + PHASES + ["Notes"]
    for col, value in enumerate(header, start=1):
        cell = sheet.cell(header_row, col, value)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="343A40")
        cell.alignment = Alignment(horizontal="center" if col > 3 else "left", wrap_text=True)

    category_fill = PatternFill("solid", fgColor="DEE2E6")
    for offset, (code, title, types, marks) in enumerate(MATRIX, start=1):
        row = header_row + offset
        code_cell = sheet.cell(row, 1, code)
        code_cell.number_format = "@"  # text, so 2.10 stays 2.10
        sheet.cell(row, 2, title)
        sheet.cell(row, 3, types)
        for index, mark in enumerate(marks):
            cell = sheet.cell(row, 4 + index, "Y" if mark == "Y" else None)
            cell.alignment = Alignment(horizontal="center")
        if not marks:
            for col in range(1, len(header) + 1):
                sheet.cell(row, col).fill = category_fill
                sheet.cell(row, col).font = Font(bold=True)

    last_row = header_row + len(MATRIX)
    validation = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    sheet.add_data_validation(validation)
    validation.add(f"D{header_row + 1}:{chr(ord('D') + len(PHASES) - 1)}{last_row}")
    for col, width in zip("ABCDEFGH", (8, 36, 18, 14, 14, 18, 14, 30)):
        sheet.column_dimensions[col].width = width
    sheet.freeze_panes = sheet.cell(header_row + 1, 4)

    notes = workbook.create_sheet("How to use")
    for line in [
        "How this matrix is read by FileCrawler",
        "",
        "- The header row needs a 'Code' column. Title and Allowed Types are optional.",
        "- Every other column with Y/N marks is a phase. Y, Yes, X, R or 1 = required; blank, N or - = not.",
        "- Codes must be numeric (1.1, 2.11). Format the Code column as Text: as numbers, 2.10 becomes 2.1.",
        "- A row whose code has no dot (e.g. '2') and no Y marks is a category heading.",
        "- Allowed Types: e.g. '.pdf; .docx'. Leave blank to accept any file type.",
        "- The submission's folders must START with the code, e.g. '2.10 Design Risk Register'.",
        "- Rows and columns above the header (titles, notes) are ignored.",
    ]:
        notes.append([line])
    notes["A1"].font = Font(bold=True, size=13)
    notes.column_dimensions["A"].width = 110
    workbook.save(path)


def write_csv(path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Code", "Title", "Allowed Types"] + PHASES)
        for code, title, types, marks in MATRIX:
            writer.writerow([code, title, types] + ["Y" if m == "Y" else "" for m in marks.ljust(len(PHASES))])


def write_zip(path: Path) -> None:
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for rel, content in sorted(SUBMISSION.items()):
            info = zipfile.ZipInfo(f"{ROOT}/{rel}", date_time=(2026, 1, 1, 0, 0, 0))
            info.external_attr = (0o40755 << 16) | 0x10 if content is None else 0o644 << 16
            archive.writestr(info, b"" if content is None else content)


def write_folder(target: Path) -> Path:
    root = target / ROOT
    for rel, content in SUBMISSION.items():
        path = root / rel
        if content is None:
            path.mkdir(parents=True, exist_ok=True)
        else:
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(content)
    return root


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--folder", type=Path, help="also unpack the sample submission into this folder")
    args = parser.parse_args()
    write_xlsx(HERE / "example_matrix.xlsx")
    write_csv(HERE / "example_matrix.csv")
    write_zip(HERE / "sample_submission.zip")
    print(f"Wrote example_matrix.xlsx, example_matrix.csv and sample_submission.zip to {HERE}")
    if args.folder:
        print(f"Sample submission unpacked to {write_folder(args.folder)}")


if __name__ == "__main__":
    main()
