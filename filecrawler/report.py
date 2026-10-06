"""Ordering, plain-language summary and export of check results."""
from __future__ import annotations

import csv
from collections import defaultdict
from pathlib import Path
from typing import Dict, List

from .models import Bucket, CheckMode, CheckReport, CheckResult, Matrix, Status

BUCKET_COLOURS = {  # (background, text) - shared by the UI and the Excel export
    Bucket.MISSING: ("F8D7DA", "842029"),
    Bucket.ATTENTION: ("FFF3CD", "664D03"),
    Bucket.PRESENT: ("D1E7DD", "0F5132"),
    Bucket.INFO: ("F1F3F5", "6C757D"),
}

EXPORT_COLUMNS = [
    "Status", "Group", "Required this phase", "Code", "Title", "Allowed types",
    "Folder(s)", "Files found", "Other files", "Notes", "AI findings",
]


def sort_results(results: List[CheckResult]) -> List[CheckResult]:
    """Missing first, then needs attention, then present, then not required."""
    return sorted(results, key=lambda r: (r.bucket.value, r.requirement.key))


def _codes(results: List[CheckResult], limit: int = 12) -> str:
    codes = [r.requirement.code for r in results]
    if len(codes) > limit:
        return ", ".join(codes[:limit]) + f" (+{len(codes) - limit} more)"
    return ", ".join(codes)


def build_evaluation(report: CheckReport, matrix: Matrix) -> List[str]:
    """Rule-based bullet points - available in every mode, no AI involved."""
    required = report.required_results
    if not required:
        return [f"Nothing is marked as required for '{report.phase}' in the matrix."]

    by_bucket: Dict[Bucket, List[CheckResult]] = defaultdict(list)
    for result in required:
        by_bucket[result.bucket].append(result)
    present = len(by_bucket[Bucket.PRESENT])
    bullets = [
        f"{present} of {len(required)} documents required for {report.phase} are present "
        f"({round(100 * present / len(required))}%)."
    ]

    missing = by_bucket[Bucket.MISSING]
    if missing:
        empty = [r for r in missing if r.status is Status.EMPTY]
        detail = f" - {len(empty)} of these have a folder but no usable file" if empty else ""
        bullets.append(f"{len(missing)} missing: {_codes(missing)}{detail}.")
    for status, text in (
        (Status.WRONG_TYPE, "in the wrong file type"),
        (Status.MISFILED, "found outside their folder - move them into the right folder"),
        (Status.INCOMPLETE, "need changes according to the AI check"),
    ):
        group = [r for r in required if r.status is status]
        if group:
            bullets.append(f"{len(group)} {text}: {_codes(group)}.")

    gaps: Dict[int, int] = defaultdict(int)
    totals: Dict[int, int] = defaultdict(int)
    for result in required:
        totals[result.requirement.category] += 1
        if result.bucket is not Bucket.PRESENT:
            gaps[result.requirement.category] += 1
    if gaps:
        worst = max(gaps, key=lambda c: (gaps[c] / totals[c], gaps[c]))
        bullets.append(
            f"Weakest area: {matrix.category_label(worst)}, where {gaps[worst]} of {totals[worst]} "
            "required documents need action."
        )
        complete = [matrix.category_label(c) for c in sorted(totals) if not gaps.get(c)]
        if complete:
            bullets.append(f"Complete categories: {'; '.join(complete)}.")
    else:
        bullets.append("Every required document is present.")

    extra = [r for r in report.results if not r.required and (r.files or r.other_files or r.misfiled_files)]
    if extra:
        bullets.append(f"{len(extra)} document(s) provided ahead of time (not required this phase): {_codes(extra)}.")
    if report.unrecognised_folders:
        shown = report.unrecognised_folders[:5]
        more = f" (+{len(report.unrecognised_folders) - 5} more)" if len(report.unrecognised_folders) > 5 else ""
        bullets.append(
            f"{len(report.unrecognised_folders)} folder(s) don't match any code in the matrix: "
            + "; ".join(shown) + more + "."
        )
    if report.mode is CheckMode.PRESENCE:
        bullets.append("Presence check only - file contents were not opened, so 'Present' does not mean correct.")
    return bullets


def _row(result: CheckResult) -> List[str]:
    return [
        result.status.label,
        result.bucket.label,
        "Yes" if result.required else "No",
        result.requirement.code,
        result.requirement.title,
        ", ".join(result.allowed_types) or "Any",
        "; ".join(result.folders),
        "; ".join(f.rel_path for f in result.files + result.misfiled_files),
        "; ".join(f.rel_path for f in result.other_files + result.empty_files),
        "; ".join(result.notes),
        "; ".join(result.ai_findings),
    ]


def export_report(report: CheckReport, path: "str | Path") -> Path:
    path = Path(path)
    if path.suffix.lower() == ".csv":
        _export_csv(report, path)
    else:
        _export_xlsx(report, path)
    return path


def _export_csv(report: CheckReport, path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(EXPORT_COLUMNS)
        for result in report.results:
            writer.writerow(_row(result))


def _export_xlsx(report: CheckReport, path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    meta = [
        ("Phase", report.phase),
        ("Check mode", report.mode.label),
        ("Submission", report.source),
        ("Matrix", report.matrix_source),
        ("Generated", report.generated_at.strftime("%Y-%m-%d %H:%M")),
        ("Missing", report.count(Bucket.MISSING)),
        ("Needs attention", report.count(Bucket.ATTENTION)),
        ("Present", report.count(Bucket.PRESENT)),
    ]
    for label, value in meta:
        summary.append([label, value])
        summary.cell(summary.max_row, 1).font = Font(bold=True)
    summary.append([])
    summary.append(["Evaluation"])
    summary.cell(summary.max_row, 1).font = Font(bold=True)
    for bullet in report.evaluation:
        summary.append(["", "• " + bullet])
    if report.ai_evaluation or report.ai_status:
        summary.append([])
        summary.append(["AI evaluation"])
        summary.cell(summary.max_row, 1).font = Font(bold=True)
        for bullet in report.ai_evaluation or [report.ai_status]:
            summary.append(["", "• " + bullet])
    if report.warnings:
        summary.append([])
        summary.append(["Warnings"])
        summary.cell(summary.max_row, 1).font = Font(bold=True)
        for warning in report.warnings:
            summary.append(["", warning])
    summary.column_dimensions["A"].width = 18
    summary.column_dimensions["B"].width = 120

    sheet = workbook.create_sheet("Results")
    sheet.append(EXPORT_COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="343A40")
    for result in report.results:
        sheet.append(_row(result))
        background, text = BUCKET_COLOURS[result.bucket]
        fill = PatternFill("solid", fgColor=background)
        for cell in sheet[sheet.max_row]:
            cell.fill = fill
            cell.font = Font(color=text)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    widths = [16, 16, 10, 8, 36, 14, 34, 40, 30, 50, 50]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    workbook.save(path)
