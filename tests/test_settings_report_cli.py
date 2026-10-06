import csv
import os

import pytest
from conftest import EXAMPLES, make_tree, matrix_from

from filecrawler.__main__ import main
from filecrawler.checker import run_check
from filecrawler.matrix import load_matrix
from filecrawler.models import CheckMode
from filecrawler.report import export_report
from filecrawler.settings import Settings, parse_type_rules


def test_type_rule_longest_prefix_wins():
    settings = Settings(type_rules={"2": [".pdf"], "2.2": [".dwg"], "2.20": [".xlsx"]})
    assert settings.allowed_types_for((2, 1)) == (".pdf",)
    assert settings.allowed_types_for((2, 2)) == (".dwg",)
    assert settings.allowed_types_for((2, 20)) == (".xlsx",)
    assert settings.allowed_types_for((3, 1)) == ()


def test_parse_type_rules():
    rules, errors = parse_type_rules("# comment\n2 = pdf\n2.2: .pdf, .DWG\n\nbad line\n3 =\n")
    assert rules == {"2": [".pdf"], "2.2": [".pdf", ".dwg"]}
    assert len(errors) == 2


def test_ignore_patterns_case_insensitive():
    settings = Settings()
    assert settings.is_ignored_file("THUMBS.DB")
    assert settings.is_ignored_file("~$Report.docx")
    assert not settings.is_ignored_file("Report.docx")


def test_settings_round_trip_and_bad_files(tmp_path):
    path = tmp_path / "s.json"
    Settings(type_rules={"1": [".pdf"]}, last_phase="P2").save(path)
    loaded = Settings.load(path)
    assert loaded.type_rules == {"1": [".pdf"]} and loaded.last_phase == "P2"

    path.write_text("{not json")
    assert Settings.load(path) == Settings()
    path.write_text('{"zero_byte_is_empty": "yes", "unknown": 1}')
    assert Settings.load(path).zero_byte_is_empty is True


@pytest.fixture
def report(settings):
    matrix = load_matrix(EXAMPLES / "example_matrix.xlsx")
    return run_check(matrix, "Phase 2 - Design", EXAMPLES / "sample_submission.zip", CheckMode.PRESENCE, settings)


def test_export_xlsx(report, tmp_path):
    from openpyxl import load_workbook
    path = export_report(report, tmp_path / "out.xlsx")
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Summary", "Results"]
    results = workbook["Results"]
    assert results.max_row == len(report.results) + 1
    assert results["A2"].value == "Empty folder"
    assert results["A2"].fill.fgColor.rgb.endswith("F8D7DA")


def test_export_csv(report, tmp_path):
    path = export_report(report, tmp_path / "out.csv")
    with path.open(encoding="utf-8-sig") as handle:
        rows = list(csv.reader(handle))
    assert rows[0][0] == "Status"
    assert len(rows) == len(report.results) + 1


def test_evaluation_bullets(report):
    text = " ".join(report.evaluation)
    assert "17 of 25" in text
    assert "5 missing: 1.3, 2.5, 2.6, 3.2, 6.4" in text
    assert "Weakest area: 2 - Design" in text
    assert "Presence check only" in text


def test_evaluation_all_present(tmp_path, settings):
    matrix = matrix_from([["1.1", "A", "", "Y", ""]])
    make_tree(tmp_path, {"1.1 A/a.pdf": b"x"})
    report = run_check(matrix, "P1", tmp_path, CheckMode.PRESENCE, settings)
    assert "Every required document is present." in report.evaluation
    report = run_check(matrix, "P2", tmp_path, CheckMode.PRESENCE, settings)
    assert report.evaluation == ["Nothing is marked as required for 'P2' in the matrix."]


def test_cli_exit_codes(tmp_path, capsys):
    matrix = str(EXAMPLES / "example_matrix.xlsx")
    assert main(["--matrix", matrix, "--phase", "design", "--source", str(EXAMPLES / "sample_submission.zip")]) == 1
    assert "== MISSING (5) ==" in capsys.readouterr().out
    assert main(["--matrix", str(tmp_path / "nope.xlsx"), "--source", str(tmp_path)]) == 2

    csv_matrix = tmp_path / "m.csv"
    csv_matrix.write_text("Code,Title,Phase 1\n1.1,Form,Y\n")
    make_tree(tmp_path / "sub", {"1.1 Form/form.pdf": b"x"})
    out = tmp_path / "out.xlsx"
    assert main(["--matrix", str(csv_matrix), "--source", str(tmp_path / "sub"), "--export", str(out)]) == 0
    assert os.path.getsize(out) > 0
